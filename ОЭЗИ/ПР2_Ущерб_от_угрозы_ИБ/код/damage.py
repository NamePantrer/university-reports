#!/usr/bin/env python3
"""Чистая прибыль и упущенная выгода при реализации угрозы ИБ (ПР2 ОЭЗИ).

Формулы методички:
  (4) ТФ  = ТФР - Тпр                       фактическое время, ч
  (3) Q   = Тр * ТФ                         объём выпуска, изд.
  (2) В   = Q * Ц                           выручка, руб.
  (1) ЧП  = В - Сб - Ур - Кр - Шт - Пр - Нл
Допущения отчёта: себестоимость пропорциональна выпуску (Сб' = Сб * Q'/Q),
Ур и Кр постоянны, Нл = 20 % от прибыли до налога.
Упущенная выгода УВ = ЧП(план) - ЧП(сценарий).

    python3 damage.py [input.json] [--xlsx книга.xlsx]
"""
import argparse
import json
from pathlib import Path


def scenario(d, downtime, fines, q_plan=None):
    tf = d["TFR"] - downtime
    q = d["Tr"] * tf
    rev = q * d["price"]
    cost = d["cost_plan"] if q_plan is None else d["cost_plan"] * q / q_plan
    ebt = rev - cost - d["admin"] - d["commercial"] - fines - d["other"]
    tax = d["tax_rate"] * ebt if ebt > 0 else 0.0
    return {"Тпр, ч": downtime, "ТФ, ч": tf, "Q, изд.": q, "Выручка, руб.": rev,
            "Себестоимость, руб.": cost, "Штрафы, руб.": fines,
            "Прибыль до налога, руб.": ebt, "Нл, руб.": tax, "Чистая прибыль, руб.": ebt - tax}


def calculate(d):
    names = list(d["scenarios"])
    plan_name = names[0]
    p = d["scenarios"][plan_name]
    plan = scenario(d, p["downtime"], p["fines"])
    res = {plan_name: plan}
    for n in names[1:]:
        s = d["scenarios"][n]
        res[n] = scenario(d, s["downtime"], s["fines"], q_plan=plan["Q, изд."])
    for n, r in res.items():
        r["Упущенная выгода, руб."] = plan["Чистая прибыль, руб."] - r["Чистая прибыль, руб."]
    return res


def fmt(x):
    return f"{x:,.0f}".replace(",", " ")


def write_xlsx(d, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "ПР2_ущерб"
    bold = Font(bold=True)
    ws["A1"] = "ПР2. Чистая прибыль и упущенная выгода (формулы (1)-(4))"; ws["A1"].font = bold
    params = [("Тр, изд./час", d["Tr"]), ("ТФР, ч", d["TFR"]), ("Ц, руб.", d["price"]),
              ("Сб (план), руб.", d["cost_plan"]), ("Ур, руб.", d["admin"]),
              ("Кр, руб.", d["commercial"]), ("Пр, руб.", d["other"]), ("Ставка налога", d["tax_rate"])]
    ref = {}
    for i, (k, v) in enumerate(params, start=3):
        ws.cell(i, 1, k); ws.cell(i, 2, v); ref[k] = f"$B${i}"
    top = 3 + len(params) + 1
    names = list(d["scenarios"])
    ws.cell(top, 1, "Показатель").font = bold
    for j, n in enumerate(names, start=2):
        ws.cell(top, j, n).font = bold
    labels = ["Тпр, ч", "Шт, руб.", "ТФ, ч", "Q, изд.", "Выручка, руб.", "Сб, руб.",
              "Прибыль до налога, руб.", "Нл, руб.", "ЧП, руб.", "Упущенная выгода, руб."]
    row = {lab: top + 1 + i for i, lab in enumerate(labels)}
    for lab, r in row.items():
        ws.cell(r, 1, lab)
    for j, n in enumerate(names, start=2):
        c = ws.cell(1, j).column_letter
        s = d["scenarios"][n]
        R = lambda lab: f"{c}{row[lab]}"  # noqa: E731
        P = lambda lab: f"$B${row[lab]}"  # noqa: E731 — столбец плана
        ws[R("Тпр, ч")] = s["downtime"]
        ws[R("Шт, руб.")] = s["fines"]
        ws[R("ТФ, ч")] = f"={ref['ТФР, ч']}-{R('Тпр, ч')}"
        ws[R("Q, изд.")] = f"={ref['Тр, изд./час']}*{R('ТФ, ч')}"
        ws[R("Выручка, руб.")] = f"={R('Q, изд.')}*{ref['Ц, руб.']}"
        ws[R("Сб, руб.")] = f"={ref['Сб (план), руб.']}*{R('Q, изд.')}/{P('Q, изд.')}"
        ws[R("Прибыль до налога, руб.")] = (f"={R('Выручка, руб.')}-{R('Сб, руб.')}-{ref['Ур, руб.']}"
                                            f"-{ref['Кр, руб.']}-{R('Шт, руб.')}-{ref['Пр, руб.']}")
        ws[R("Нл, руб.")] = f"=MAX(0,{ref['Ставка налога']}*{R('Прибыль до налога, руб.')})"
        ws[R("ЧП, руб.")] = f"={R('Прибыль до налога, руб.')}-{R('Нл, руб.')}"
        ws[R("Упущенная выгода, руб.")] = f"={P('ЧП, руб.')}-{R('ЧП, руб.')}"
    ws.column_dimensions["A"].width = 28
    for c in "BCD":
        ws.column_dimensions[c].width = 18
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", nargs="?", default=Path(__file__).with_name("input.json"))
    ap.add_argument("--xlsx")
    a = ap.parse_args()
    d = json.loads(Path(a.input).read_text(encoding="utf-8"))
    res = calculate(d)
    names = list(res)
    print(d["enterprise"])
    print(f"{'Показатель':26s}" + "".join(f"{n:>18s}" for n in names))
    for key in res[names[0]]:
        print(f"{key:26s}" + "".join(f"{fmt(res[n][key]):>18s}" for n in names))
    if len(names) >= 3:
        saved = res[names[1]]["Упущенная выгода, руб."] - res[names[2]]["Упущенная выгода, руб."]
        print(f"\nСЗИ снижают упущенную выгоду на {fmt(saved)} руб. на один инцидент: "
              f"вложение окупается, если контур защиты стоит дешевле.")
    if a.xlsx:
        write_xlsx(d, a.xlsx)
        print(f"Книга с формулами: {a.xlsx}")


if __name__ == "__main__":
    main()
