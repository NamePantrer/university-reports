#!/usr/bin/env python3
"""Оценка экономической эффективности программного проекта (ПР4 ОЭЗИ).

Статические методы (Павлова Е.А., п. 4.2):
  Qk  = F / (P - V)                       точка безубыточности (4.1)
  RI  = (1/T * sum P_t) / I               рентабельность инвестиций (4.3)
  ARR = (sum P_t / T) / (0.5 (I + L))     коэффициент эффективности (4.4)
  PP  — простой срок окупаемости
Динамические методы (п. 4.4):
  NPV = -I + sum NCF_t / (1+R)^t          (4.16)
  A   = NPV * R (1+R)^T / ((1+R)^T - 1)   аннуитет (4.18)
  PI  = sum PV / I                        (4.21)
  IRR — ставка, при которой NPV = 0       (4.22), ищется бисекцией
  DPP — дисконтированный срок окупаемости
В качестве годовой прибыли P_t, как и в отчёте, взят NCF_t = CIF_t - COF_t.

    python3 efficiency.py [input.json] [--xlsx книга.xlsx]
"""
import argparse
import json
from pathlib import Path


def npv(rate, I, ncf):
    return -I + sum(c / (1 + rate) ** t for t, c in enumerate(ncf, 1))


def irr(I, ncf, lo=-0.99, hi=10.0, eps=1e-10):
    f_lo = npv(lo, I, ncf)
    if f_lo * npv(hi, I, ncf) > 0:
        raise ValueError("IRR вне диапазона поиска")
    for _ in range(300):
        mid = (lo + hi) / 2
        f_mid = npv(mid, I, ncf)
        if abs(f_mid) < eps:
            break
        if f_lo * f_mid < 0:
            hi = mid
        else:
            lo, f_lo = mid, f_mid
    return mid


def payback(I, flows):
    acc = 0.0
    for t, c in enumerate(flows, 1):
        if acc + c >= I:
            return t - 1 + (I - acc) / c
        acc += c
    return float("inf")


def calculate(d):
    ncf = [a - b for a, b in zip(d["CIF"], d["COF"])]
    T, I, R, L = len(ncf), d["I"], d["R"], d["L"]
    pv = [c / (1 + R) ** t for t, c in enumerate(ncf, 1)]
    avg = sum(ncf) / T
    res = {
        "NCF": ncf, "PV": pv,
        "Qk": d["F"] / (d["P"] - d["V"]),
        "RI": avg / I,
        "ARR": avg / (0.5 * (I + L)),
        "PP": payback(I, ncf),
        "NPV": sum(pv) - I,
        "PI": sum(pv) / I,
        "IRR": irr(I, ncf),
        "DPP": payback(I, pv),
    }
    k = R * (1 + R) ** T / ((1 + R) ** T - 1)
    res["annuity_k"] = k
    res["A"] = res["NPV"] * k
    return res


def fmt(x):
    return f"{x:,.0f}".replace(",", " ")


def write_xlsx(d, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(); ws = wb.active; ws.title = "ПР4"
    b = Font(bold=True)
    ws["A1"] = "ПР4. Эффективность ИС «АудитИБ» (формулы)"; ws["A1"].font = b
    params = [("I, руб.", d["I"]), ("R", d["R"]), ("L, руб.", d["L"]), ("P, руб./клиент", d["P"]),
              ("V, руб./клиент", d["V"]), ("F, руб./год", d["F"]), ("Спрос Q, клиентов", d["demand"])]
    ref = {}
    for i, (k, v) in enumerate(params, 3):
        ws.cell(i, 1, k); ws.cell(i, 2, v); ref[k] = f"$B${i}"
    h = 3 + len(params) + 1
    for j, t in enumerate(["Год t", "CIF", "COF", "NCF", "α_t", "PV", "Нараст. NCF", "Нараст. PV"], 1):
        ws.cell(h, j, t).font = b
    r0 = h + 1
    ws.cell(r0, 1, 0); ws.cell(r0, 4, f"=-{ref['I, руб.']}"); ws.cell(r0, 6, f"=D{r0}")
    ws.cell(r0, 7, f"=D{r0}"); ws.cell(r0, 8, f"=F{r0}")
    T = len(d["CIF"])
    for t in range(1, T + 1):
        r = r0 + t
        ws.cell(r, 1, t); ws.cell(r, 2, d["CIF"][t - 1]); ws.cell(r, 3, d["COF"][t - 1])
        ws.cell(r, 4, f"=B{r}-C{r}")
        ws.cell(r, 5, f"=1/(1+{ref['R']})^A{r}")
        ws.cell(r, 6, f"=D{r}*E{r}")
        ws.cell(r, 7, f"=G{r - 1}+D{r}")
        ws.cell(r, 8, f"=H{r - 1}+F{r}")
    first, last = r0 + 1, r0 + T
    s = last + 2
    ws.cell(s, 1, "Показатель").font = b; ws.cell(s, 2, "Значение").font = b
    rows = [
        ("Qk = F/(P−V)", f"={ref['F, руб./год']}/({ref['P, руб./клиент']}-{ref['V, руб./клиент']})"),
        ("RI", f"=AVERAGE(D{first}:D{last})/{ref['I, руб.']}"),
        ("ARR", f"=AVERAGE(D{first}:D{last})/(0.5*({ref['I, руб.']}+{ref['L, руб.']}))"),
        ("NPV", f"=NPV({ref['R']},D{first}:D{last})-{ref['I, руб.']}"),
        ("PI", f"=SUM(F{first}:F{last})/{ref['I, руб.']}"),
        ("IRR", f"=IRR(D{r0}:D{last})"),
        ("Аннуитет A", f"=B{s + 4}*{ref['R']}*(1+{ref['R']})^{T}/((1+{ref['R']})^{T}-1)"),
        ("PP, лет", f"=SUMPRODUCT(--(G{first}:G{last}<0))+"
                    f"ABS(INDEX(G{r0}:G{last},SUMPRODUCT(--(G{first}:G{last}<0))+1))/"
                    f"INDEX(D{first}:D{last},SUMPRODUCT(--(G{first}:G{last}<0))+1)"),
        ("DPP, лет", f"=SUMPRODUCT(--(H{first}:H{last}<0))+"
                     f"ABS(INDEX(H{r0}:H{last},SUMPRODUCT(--(H{first}:H{last}<0))+1))/"
                     f"INDEX(F{first}:F{last},SUMPRODUCT(--(H{first}:H{last}<0))+1)"),
    ]
    for i, (k, f) in enumerate(rows, s + 1):
        ws.cell(i, 1, k); ws.cell(i, 2, f)
    ws.column_dimensions["A"].width = 22
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 14
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", nargs="?", default=Path(__file__).with_name("input.json"))
    ap.add_argument("--xlsx")
    a = ap.parse_args()
    d = json.loads(Path(a.input).read_text(encoding="utf-8"))
    r = calculate(d)
    R = d["R"]
    print(d["project"])
    print(f"{'t':>2s} {'CIF':>10s} {'COF':>10s} {'NCF':>10s} {'α_t':>7s} {'PV':>10s}")
    for t, (ci, co, n, p) in enumerate(zip(d["CIF"], d["COF"], r["NCF"], r["PV"]), 1):
        print(f"{t:>2d} {fmt(ci):>10s} {fmt(co):>10s} {fmt(n):>10s} {1 / (1 + R) ** t:7.4f} {fmt(p):>10s}")
    yes = lambda c: "да" if c else "НЕТ"  # noqa: E731
    T = len(r["NCF"])
    print("\nПоказатель   Значение          Критерий           Итог")
    print(f"Qk           {r['Qk']:.2f} клиента      Q={d['demand']} > Qk          {yes(d['demand'] > r['Qk'])}")
    print(f"RI           {r['RI'] * 100:.2f} %          >= R={R * 100:.0f} %         {yes(r['RI'] >= R)}")
    print(f"ARR          {r['ARR'] * 100:.2f} %          ориентир           —")
    print(f"PP           {r['PP']:.2f} года         <= 3               {yes(r['PP'] <= 3)}")
    print(f"NPV          {fmt(r['NPV'])} руб.      > 0                {yes(r['NPV'] > 0)}")
    print(f"A            {fmt(r['A'])} руб./год    справочно (k={r['annuity_k']:.6f})")
    print(f"PI           {r['PI']:.3f}             > 1                {yes(r['PI'] > 1)}")
    print(f"IRR          {r['IRR'] * 100:.2f} %          > {R * 100:.0f} %             {yes(r['IRR'] > R)}")
    print(f"DPP          {r['DPP']:.2f} года         <= {T}               {yes(r['DPP'] <= T)}")
    if a.xlsx:
        write_xlsx(d, a.xlsx)
        print(f"\nКнига с формулами: {a.xlsx}")


if __name__ == "__main__":
    main()
