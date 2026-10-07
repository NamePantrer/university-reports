#!/usr/bin/env python3
"""Четыре способа ценообразования информационного продукта (ПР1 ОЭЗИ).

  1. Затратный:          (затраты + затраты * норма прибыли) / копии
  2. По времени автора:  (месячный доход * часы_продукта / фонд_времени) / копии
  3. По продажам аналога: цена_аналога * (1 + средний месячный прирост продаж)
  4. Доля от полезности: потенциальный доход покупателя * авторский процент

    python3 pricing.py [input.json]             # расчёт, как в отчёте
    python3 pricing.py --xlsx расчет.xlsx       # книга Excel с живыми формулами
"""
import argparse
import json
from pathlib import Path


def cost_based(costs, copies, rate):
    total = sum(costs.values())
    return (total + total * rate) / copies


def time_based(income, fund, hours, copies):
    share = round(hours / fund, 4)          # в методичке доля округляется до 4 знаков
    return income * share / copies, share


def sales_based(sales, analog_price):
    growth = [(b - a) / a for a, b in zip(sales, sales[1:])]
    avg = round(sum(growth) / len(growth), 4)
    return analog_price * (1 + avg), avg


def value_based(benefit, percent):
    return benefit * percent


def calculate(d):
    p1 = cost_based(d["costs"], d["copies"], d["profit_rate"])
    p2, share = time_based(d["monthly_income"], d["work_fund_hours"], d["product_hours"], d["copies"])
    p3, growth = sales_based(d["analog_sales"], d["analog_price"])
    p4 = value_based(d["buyer_benefit"], d["author_percent"])
    return {"total_costs": sum(d["costs"].values()), "share": share, "growth": growth,
            "prices": [round(p1), round(p2), round(p3), round(p4)]}


def write_xlsx(d, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "ПР1_цена"
    b = Font(bold=True)
    ws["A1"] = "ПР1. Цена курса — четыре способа (формулы)"; ws["A1"].font = b
    ws.append([])
    ws.append(["Статья затрат", "Сумма, руб."])
    first = ws.max_row + 1
    for k, v in d["costs"].items():
        ws.append([k, v])
    last = ws.max_row
    ws.append(["Итого затраты", f"=SUM(B{first}:B{last})"]); tot = ws.max_row
    ws.append([])
    rows = {}
    for key, label, val in [
        ("copies", "Расчетное количество копий, шт.", d["copies"]),
        ("rate", "Норма прибыли", d["profit_rate"]),
        ("income", "Желаемый месячный доход, руб.", d["monthly_income"]),
        ("fund", "Фонд рабочего времени, ч", d["work_fund_hours"]),
        ("hours", "Время на продукт, ч", d["product_hours"]),
        ("analog", "Цена аналога, руб.", d["analog_price"]),
        ("benefit", "Потенциальный доход покупателя, руб.", d["buyer_benefit"]),
        ("pct", "Авторский процент", d["author_percent"]),
    ]:
        ws.append([label, val]); rows[key] = f"B{ws.max_row}"
    ws.append(["Продажи аналога по месяцам, ед."] + d["analog_sales"]); srow = ws.max_row
    n = len(d["analog_sales"])
    cols = "BCDEFGHIJ"
    growth_terms = "+".join(f"({cols[i + 1]}{srow}-{cols[i]}{srow})/{cols[i]}{srow}" for i in range(n - 1))
    ws.append([])
    ws.append(["Расчет", "Значение"]); ws.cell(ws.max_row, 1).font = b
    ws.append(["Способ 1: процент к затратам, руб.", f"=ROUND((B{tot}+B{tot}*{rows['rate']})/{rows['copies']},0)"])
    ws.append(["Доля времени", f"=ROUND({rows['hours']}/{rows['fund']},4)"]); sh = ws.max_row
    ws.append(["Способ 2: время автора, руб.", f"=ROUND({rows['income']}*B{sh}/{rows['copies']},0)"])
    ws.append(["Средний месячный прирост", f"=ROUND(({growth_terms})/{n - 1},4)"]); gr = ws.max_row
    ws.append(["Способ 3: продажи аналога, руб.", f"=ROUND({rows['analog']}*(1+B{gr}),0)"])
    ws.append(["Способ 4: доля от полезности, руб.", f"=ROUND({rows['benefit']}*{rows['pct']},0)"])
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input", nargs="?", default=Path(__file__).with_name("input.json"))
    ap.add_argument("--xlsx")
    a = ap.parse_args()
    d = json.loads(Path(a.input).read_text(encoding="utf-8"))
    r = calculate(d)
    print(d["product"])
    total = f"{r['total_costs']:,}".replace(",", " ")
    print(f"Затраты: {total} руб., копий: {d['copies']}")
    names = ["процент к затратам", "время автора", "продажи аналога", "доля от полезности"]
    for i, (n, p) in enumerate(zip(names, r["prices"]), 1):
        print(f"  Способ {i} ({n}): {p} руб.")
    print(f"  доля времени = {r['share']:.4f}; средний прирост продаж = {r['growth']:.4f}")
    lo, hi = min(r["prices"]), max(r["prices"])
    print(f"Разброс: {lo}-{hi} руб. ({hi / lo:.1f} раза). Рабочая розница (затратный метод): {r['prices'][0]} руб.")
    if a.xlsx:
        write_xlsx(d, a.xlsx)
        print(f"Книга с формулами: {a.xlsx}")


if __name__ == "__main__":
    main()
